"""Merge many Excel/CSV files into one workbook, tagging each row with its source file."""
import csv
from pathlib import Path

from openpyxl import Workbook, load_workbook


def read_rows(path: Path) -> list[list]:
    if path.suffix.lower() == ".csv":
        with path.open(newline="", encoding="utf-8-sig") as f:
            return [row for row in csv.reader(f)]
    wb = load_workbook(path, read_only=True, data_only=True)
    ws = wb.active
    rows = [list(r) for r in ws.iter_rows(values_only=True)]
    wb.close()
    return rows


def merge(inputs: list[Path], output: Path, add_source: bool = True) -> int:
    """Merge files that share a header row. Columns are unioned by name. Returns rows written."""
    headers: list[str] = []
    records: list[dict] = []
    for path in inputs:
        rows = read_rows(path)
        if not rows:
            continue
        head = [str(h).strip() if h is not None else "" for h in rows[0]]
        for h in head:
            if h and h not in headers:
                headers.append(h)
        for row in rows[1:]:
            if all(v in (None, "") for v in row):
                continue
            rec = {head[i]: v for i, v in enumerate(row) if i < len(head) and head[i]}
            rec["source_file"] = path.name
            records.append(rec)

    cols = headers + (["source_file"] if add_source else [])
    wb = Workbook()
    ws = wb.active
    ws.title = "Merged"
    ws.append(cols)
    for rec in records:
        ws.append([rec.get(c) for c in cols])
    wb.save(output)
    return len(records)


def add_parser(sub):
    p = sub.add_parser("excel-merge", help="Merge .xlsx/.csv files into one workbook")
    p.add_argument("inputs", nargs="+", type=Path)
    p.add_argument("-o", "--output", type=Path, default=Path("merged.xlsx"))
    p.add_argument("--no-source", action="store_true", help="Don't add a source_file column")
    p.set_defaults(func=lambda a: print(
        f"Wrote {merge(a.inputs, a.output, not a.no_source)} rows to {a.output}"))
