import csv

import pytest
from openpyxl import Workbook, load_workbook
from pypdf import PdfReader, PdfWriter

from officekit import bulk_rename, clean_csv, excel_merge, pdf_tools, scrape_table


def make_pdf(path, pages):
    w = PdfWriter()
    for _ in range(pages):
        w.add_blank_page(width=200, height=200)
    with path.open("wb") as f:
        w.write(f)


def test_excel_merge_unions_columns(tmp_path):
    a = tmp_path / "a.csv"
    a.write_text("name,email\nAnn,ann@x.com\n")
    wb = Workbook()
    wb.active.append(["name", "phone"])
    wb.active.append(["Bob", "123"])
    wb.save(tmp_path / "b.xlsx")

    out = tmp_path / "out.xlsx"
    assert excel_merge.merge([a, tmp_path / "b.xlsx"], out) == 2
    rows = list(load_workbook(out).active.iter_rows(values_only=True))
    assert rows[0] == ("name", "email", "phone", "source_file")
    assert rows[2] == ("Bob", None, "123", "b.xlsx")


def test_pdf_merge_and_split(tmp_path):
    make_pdf(tmp_path / "a.pdf", 2)
    make_pdf(tmp_path / "b.pdf", 3)
    merged = tmp_path / "m.pdf"
    assert pdf_tools.merge([tmp_path / "a.pdf", tmp_path / "b.pdf"], merged) == 5

    outs = pdf_tools.split(merged, tmp_path / "s", "1-2,4-")
    assert [len(PdfReader(o).pages) for o in outs] == [2, 2]
    assert len(pdf_tools.split(merged, tmp_path / "each")) == 5


def test_pdf_bad_range(tmp_path):
    with pytest.raises(ValueError):
        pdf_tools.parse_ranges("3-9", 5)


def test_rename_plan_and_apply(tmp_path):
    for n in ("b.txt", "a.txt"):
        (tmp_path / n).write_text(n)
    pairs = bulk_rename.plan(list(tmp_path.iterdir()), "doc_{n:02}{ext}")
    assert [(o.name, n.name) for o, n in pairs] == [("a.txt", "doc_01.txt"), ("b.txt", "doc_02.txt")]
    bulk_rename.apply(pairs)
    assert (tmp_path / "doc_01.txt").read_text() == "a.txt"


def test_rename_rejects_duplicates(tmp_path):
    for n in ("a.txt", "b.txt"):
        (tmp_path / n).write_text("")
    with pytest.raises(ValueError):
        bulk_rename.plan(list(tmp_path.iterdir()), "same{ext}")


def test_clean_csv(tmp_path):
    src = tmp_path / "in.csv"
    src.write_text(
        "Name , Email,Phone\n"
        " Ann ,ANN@X.com,(555) 123-4567\n"
        "Ann again,ann@x.com ,555\n"
        ",,\n"
        "Bob,not-an-email,\n")
    out = tmp_path / "out.csv"
    stats = clean_csv.clean(src, out)
    assert stats == {"input": 4, "empty": 1, "duplicates": 1, "invalid_email": 1, "output": 2}
    rows = list(csv.DictReader(out.open()))
    assert rows[0] == {"Name": "Ann", "Email": "ann@x.com", "Phone": "5551234567"}


def test_parse_tables_and_links():
    html = """<table><tr><th>A</th><th>B</th></tr><tr><td>1</td><td>2</td></tr></table>
              <a href="/x">X</a>"""
    assert scrape_table.parse_tables(html) == [[["A", "B"], ["1", "2"]]]
    assert scrape_table.parse_links(html, "https://e.com/p")[1] == ["X", "https://e.com/x"]
